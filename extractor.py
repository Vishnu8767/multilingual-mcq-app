import PyPDF2
import docx
import pptx

def extract_text(uploaded_file) -> str:
    """Extracts text from PDF, DOCX, PPTX, or TXT files."""
    filename = uploaded_file.name.lower()
    text = ""
    
    try:
        if filename.endswith(".pdf"):
            pdf_reader = PyPDF2.PdfReader(uploaded_file)
            for page in pdf_reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
                    
        elif filename.endswith(".docx"):
            doc = docx.Document(uploaded_file)
            for para in doc.paragraphs:
                text += para.text + "\n"
                
        elif filename.endswith(".pptx"):
            presentation = pptx.Presentation(uploaded_file)
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
                        
        elif filename.endswith(".txt"):
            text = uploaded_file.getvalue().decode("utf-8")
            
        else:
            return ""
            
    except Exception as e:
        print(f"Error reading file: {e}")
        return ""
        
    return text.strip()